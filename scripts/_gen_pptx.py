# -*- coding: utf-8 -*-
"""生成 XH-202609 方案介绍 PPT（17 页，暗色科技风）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

BG = RGBColor(0x0A, 0x11, 0x28)
ACCENT = RGBColor(0x00, 0xFF, 0x9D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0x9A, 0xA6, 0xC0)
LIGHT = RGBColor(0xD7, 0xDE, 0xEE)
CARD = RGBColor(0x14, 0x1F, 0x3D)
CARD2 = RGBColor(0x10, 0x18, 0x30)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG


def title_bar(slide, text, subtitle=None):
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.05))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    r.font.size = Pt(30); r.font.bold = True; r.font.color.rgb = ACCENT
    if subtitle:
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = subtitle
        r2.font.size = Pt(14); r2.font.color.rgb = GREY
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.5), Inches(2.4), Inches(0.07))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def bullets(slide, items, top=1.85, left=0.8, width=11.7, height=5.2, size=18):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        text, lvl = it if isinstance(it, tuple) else (it, 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        r = p.add_run()
        r.text = ("•  " if lvl == 0 else "–  ") + text
        r.font.size = Pt(size - lvl * 2)
        r.font.color.rgb = WHITE if lvl == 0 else LIGHT
        p.space_after = Pt(9)


def big_text(slide, text, sub=None, top=2.6):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(top), Inches(11.7), Inches(2.2))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(40); r.font.bold = True; r.font.color.rgb = ACCENT
    if sub:
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(18); r2.font.color.rgb = LIGHT


def add_table(slide, data, top=2.0, left=0.8, width=11.7, col_w=None):
    rows, cols = len(data), len(data[0])
    gf = slide.shapes.add_table(rows, cols, Inches(left), Inches(top), Inches(width), Inches(0.5 * rows))
    table = gf.table
    if col_w:
        for j, w in enumerate(col_w):
            table.columns[j].width = Inches(w)
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            cell.vertical_anchor = PP_ALIGN.CENTER if False else None
            para = cell.text_frame.paragraphs[0]
            para.font.size = Pt(14)
            if i == 0:
                para.font.bold = True; para.font.color.rgb = BG
                cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
            else:
                para.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = CARD if i % 2 else CARD2


def footer(slide, n):
    tb = slide.shapes.add_textbox(Inches(11.6), Inches(7.05), Inches(1.5), Inches(0.35))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
    r = p.add_run(); r.text = f"XH-202609 · {n}"
    r.font.size = Pt(10); r.font.color.rgb = GREY


def slide():
    s = prs.slides.add_slide(BLANK)
    set_bg(s)
    return s


# ---- S1 封面 ----
s = slide()
big_text(s, "基于 LLM 与确定性规则混合架构的\n安全事件研判智能体", top=2.2)
tb = s.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11.7), Inches(1.2))
tf = tb.text_frame; tf.word_wrap = True
for line, col, sz in [
    ("安全事件初步研判与误报剔除智能体 · 赛题 XH-202609（自主决策闭环）", LIGHT, 18),
    ("挑战杯「揭榜挂帅」专项赛 · 杭州安恒信息技术股份有限公司命题", GREY, 14),
]:
    p = tf.paragraphs[0] if line.startswith("安全事件") else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = line; r.font.size = Pt(sz); r.font.color.rgb = col

# ---- S2 痛点 ----
s = slide(); title_bar(s, "安全运营的三大痛点", "为什么需要自主决策智能体")
bullets(s, [
    "告警泛滥：海量告警淹没真实攻击，分析师疲劳、漏判风险高",
    "误报率高：规则引擎误报多，人工逐条核对成本高、响应慢",
    "大模型直接研判：成本高、延迟大、存在幻觉与格式漂移",
    "核心矛盾：专家经验难以规模化复用，防御体系智能化不足",
])

# ---- S3 方案一句话 ----
s = slide(); title_bar(s, "我们的方案：规则先筛 · LLM 深判 · 人工兜底", "确定性预筛 + LLM 决策核心 + 置信度门控 + HITL 闭环")
bullets(s, [
    "确定性预筛：明确攻击/误报毫秒级裁决，不确定才交给 LLM（节省 70% 调用）",
    "LLM 决策核心：综合多源证据推理裁决，而非「数关键词」硬编码规则",
    "置信度门控 + HITL：低置信不自动处置，人工复核回写形成闭环",
])

# ---- S4 架构 ----
s = slide(); title_bar(s, "整体架构：六层自主决策闭环")
bullets(s, [
    "输入层：告警 / 日志 / 接口文档接入",
    "① 确定性预筛（毫秒级三态分流）",
    "② LLM 深度研判（DeepSeek/Qwen）+ 置信度门控 + 误报记忆 + RAG grounding",
    "③ 多 Agent 协作：Triage → Hunt → Respond → Report",
    "④ 闭环处置：处置预案 / 工单 / 隔离建议 + HITL 人工复核",
    "⑤ 审计账本（Investigation Ledger）：全证据链可回放、可导出",
])

# ---- S5 预筛器 ----
s = slide(); title_bar(s, "确定性预筛器（自主初判）", "阶段 1")
bullets(s, [
    "三态分流：AUTO_CLOSE（误报自动关闭）/ AUTO_ESCALATE（直接定级）/ NEED_LLM（深判）",
    "规则可配置：7 条攻击规则 + 4 条误报规则（data/prefilter_rules.json）",
    "效果：LLM 调用节省率 70%；EVENT-001 攻击 0.2ms 定级，EVENT-002 误报 0.1ms 关闭",
])

# ---- S6 置信门控 ----
s = slide(); title_bar(s, "置信度门控 + HITL（负责任的自主任）", "阶段 2")
bullets(s, [
    "置信度数值化：high 0.9 / medium 0.7 / low 0.5",
    "低置信 → needs_human_review=true，不自动处置、不生成高危工单",
    "Web 复核按钮 + 回写 review_feedback.jsonl，供误报记忆复用",
])

# ---- S7 多Agent ----
s = slide(); title_bar(s, "多 Agent 协作 + 知识库", "阶段 3-4")
bullets(s, [
    "Triage → Hunt → Respond → Report 顺序协作，前序输出作为后序输入",
    "每 Agent 独立可测、可 mock、失败降级（fail-open 不中断链路）",
    "RAG 知识库：21 条含 MITRE ATT&CK / CVE / 预案，prompt 注入 grounding 依据降幻觉",
])

# ---- S8 创新1 ----
s = slide(); title_bar(s, "创新 1：预筛 + LLM 混合架构")
bullets(s, [
    "市面多数方案：单模型全量研判 → 慢、贵、易幻觉",
    "我们：规则先筛，明确攻击/误报毫秒级裁决，仅灰色地带交给 LLM",
    "价值：快（毫秒级）+ 省（70% 调用）+ 准（LLM 聚焦难例）",
])

# ---- S9 创新2-3 ----
s = slide(); title_bar(s, "创新 2-3：置信门控 + 误报记忆")
bullets(s, [
    "置信度门控：低置信不自动处置，人工兜底形成闭环（HITL）",
    "误报记忆：相似历史研判注入 prompt（同主机/同进程/同行为加权）",
    "价值：同类误报不重复出现，系统从历史中持续学习",
])

# ---- S10 创新4-5 ----
s = slide(); title_bar(s, "创新 4-5：多 Agent + 审计账本")
bullets(s, [
    "多 Agent 顺序协作：初判/狩猎/处置/报告各司其职，体现自主决策分工",
    "Investigation Ledger：每次研判的 prompt、LLM 响应、工具调用、证据引用全记录",
    "价值：决策可回放、可审计、可导出，评审可逐条核查依据",
])

# ---- S11 三模式对比 ----
s = slide(); title_bar(s, "三模式对比（核心数据）", "同一评测集，真实 DeepSeek 实测")
add_table(s, [
    ["模式", "通过率", "成本", "LLM 调用"],
    ["纯规则", "低", "0", "0/10"],
    ["纯 LLM", "中（依赖模型）", "高", "10/10"],
    ["预筛+LLM 混合", "90%", "低", "3/10"],
], top=2.1, col_w=[3.5, 2.6, 2.6, 3.0])
bullets(s, [("混合模式通过率最高、成本最低——LLM 决策价值在难例上被精准调用", 0)], top=5.0)

# ---- S12 评测指标 ----
s = slide(); title_bar(s, "评测指标", "10 案例评测基准（已重标）")
bullets(s, [
    "综合通过率 90%（风险分级 90% / 误报识别 100% / 事件类型 100%）",
    "输出完整率 100%",
    "误报关闭率 100% · 成本估算 · 置信度校准（高置信 100% 正确）",
])

# ---- S13 效率与成本 ----
s = slide(); title_bar(s, "效率与成本")
bullets(s, [
    "LLM 调用节省率 70%（预筛分流确定事件，不调模型）",
    "预筛直达毫秒级 vs LLM 深判秒级",
    "Mock 模式零依赖可跑，真实 LLM 一键切换（DEEPSEEK_API_KEY）",
])

# ---- S14 基准对标 ----
s = slide(); title_bar(s, "业界基准对标与体系化创新")
bullets(s, [
    "对标：CyberSOCEval / SecRespond / ExCyTIn-Bench（见 docs/BENCHMARK.md）",
    "论证混合架构在安全研判上的合理性与先进性",
    "14 项体系化创新点（见 docs/INNOVATION.md v2.0）",
])

# ---- S15 双赛题 ----
s = slide(); title_bar(s, "双赛题覆盖：一套系统，两种论证")
bullets(s, [
    "XH-202614：LLM 是核心决策引擎（三模式对比评测 + 工具调用链）",
    "XH-202609：自主决策闭环（多 Agent + 预筛 + HITL + 记忆 + 审计）",
    "答辩主线：以 LLM 为决策引擎，构建完整自主决策闭环",
])

# ---- S16 可复现与落地 ----
s = slide(); title_bar(s, "可复现性与落地")
bullets(s, [
    "Mock 模式零依赖可跑，真实 LLM 一键切换",
    "一键启动：python -m security_agent.cli serve --port 8080",
    "落地场景：红蓝对抗研判、自动化漏洞挖掘辅助、企业 SOC 智能增强",
    "决赛就绪：支持 3 人 + Agent 人机协同形态（HITL 复核回写）",
])

# ---- S17 结尾 ----
s = slide()
big_text(s, "谢谢各位评委！", sub="Q&A · 团队负责人联系方式见申报表", top=2.8)

# 样式后处理：中文字体
for sl in prs.slides:
    for shape in sl.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.font.name = "Microsoft YaHei"
                    rPr = run._r.get_or_add_rPr()
                    ea = rPr.find(qn('a:ea'))
                    if ea is None:
                        ea = rPr.makeelement(qn('a:ea'), {})
                        rPr.append(ea)
                    ea.set('typeface', 'Microsoft YaHei')
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            run.font.name = "Microsoft YaHei"
                            rPr = run._r.get_or_add_rPr()
                            ea = rPr.find(qn('a:ea'))
                            if ea is None:
                                ea = rPr.makeelement(qn('a:ea'), {})
                                rPr.append(ea)
                            ea.set('typeface', 'Microsoft YaHei')

out = r"E:\Program\2026挑战杯：Security-Agent-安全事件研判智能体\submission\02_文档材料\方案介绍PPT.pptx"
prs.save(out)
print("SAVED", out, "SLIDES", len(prs.slides._sldIdLst))
